from docx import Document
from pathlib import Path
import re, json
p = Path(r'C:\escola\pap\code\hexomel_vite\estudo\relatorio\analisei e acreentar\relatorio_v1.docx')
d = Document(str(p))
paras=[]
for i, par in enumerate(d.paragraphs,1):
    text=' '.join(par.text.split())
    if text:
        paras.append({'i':i,'style':par.style.name if par.style else '', 'text':text})
headings=[]
for x in paras:
    t=x['text']; s=x['style']
    if 'Heading' in s or 'Titulo' in s or re.match(r'^\d+(?:\.\d+)*\.?\s+', t) or t in ['Conclusão','Referências bibliográficas','Glossário','Apêndices/Anexos']:
        headings.append(x)
placeholders=[]
patterns=['Lorem ipsum','Apagar','Capítulo…','Subcapítulo','0000','frases tipo','Exemplo de frases','eletrocardiógrafo','gerações de evolução da tecnologia']
for x in paras:
    if any(pat.lower() in x['text'].lower() for pat in patterns):
        placeholders.append(x)
short_sections=[]
# crude section lengths by heading index
for idx,h in enumerate(headings):
    nxt=headings[idx+1]['i'] if idx+1<len(headings) else 10**9
    count=sum(1 for x in paras if h['i'] < x['i'] < nxt and len(x['text'])>40)
    short_sections.append({'heading':h['text'], 'line':h['i'], 'count':count})
numbered=[h['text'] for h in headings if re.match(r'^\d', h['text'])]
print('PARAS', len(paras), 'HEADINGS', len(headings), 'PLACEHOLDERS', len(placeholders))
print('\nHEADINGS')
for h in headings: print(f"{h['i']:03d} [{h['style']}] {h['text']}")
print('\nPLACEHOLDERS')
for x in placeholders[:80]: print(f"{x['i']:03d} [{x['style']}] {x['text'][:180]}")
print('\nSHORT')
for s in short_sections:
    if s['count'] <= 1: print(f"{s['line']:03d} {s['count']} {s['heading']}")
